from datetime import date
import csv
import difflib
import io
import os
import re
import sqlite3
import uuid
import zipfile
from pathlib import Path

import cv2
import numpy as np
import pytesseract
from openpyxl import load_workbook
from dotenv import load_dotenv
from openpyxl.utils.exceptions import InvalidFileException
from flask import (
    Flask,
    jsonify,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from database import create_tables, get_db_connection, seed_demo_data

load_dotenv()

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-only-change-me")
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
if os.environ.get("TESSERACT_CMD"):
    pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]


@app.after_request
def add_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = os.environ.get("FRONTEND_URL", "http://localhost:5173")
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    response.headers["Access-Control-Allow-Credentials"] = "true"
    return response


@app.route("/api/health")
def api_health():
    return jsonify({"status": "ok", "service": "smart-attendance-api"})


@app.route("/api/session")
def api_session():
    return jsonify({
        "logged_in": bool(session.get("user_id")),
        "role": session.get("role"),
        "name": session.get("name"),
        "selected_class_id": session.get("selected_class_id"),
    })


@app.route("/api/logout", methods=["POST"])
def api_logout():
    session.clear()
    return jsonify({"success": True})


def class_access_allowed(class_id):
    if session.get("role") not in {"admin", "faculty"} or class_id is None:
        return False
    connection = get_db_connection()
    active_class = connection.execute("SELECT 1 FROM classes WHERE id = ? AND active = 1", (class_id,)).fetchone()
    connection.close()
    return active_class is not None


def detect_face_crop(image):
    cascade = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    faces = cascade.detectMultiScale(gray, scaleFactor=1.1, minNeighbors=5, minSize=(80, 80))
    if len(faces) != 1:
        return None, "Capture exactly one clear face in the frame."
    x, y, width, height = faces[0]
    return cv2.resize(gray[y : y + height, x : x + width], (200, 200)), None


def build_dashboard_payload():
    class_id = session.get("selected_class_id")
    if not class_id:
        return {"class": None, "stats": {}, "students": [], "subjects": [], "notifications": []}

    connection = get_db_connection()
    class_row = connection.execute(
        "SELECT c.id, c.year, c.department, c.division, a.academic_year FROM classes c JOIN academic_years a ON a.id = c.academic_year_id WHERE c.id = ?",
        (class_id,),
    ).fetchone()
    students = connection.execute("SELECT id, student_uid, roll_no, name, email, phone, class_id, face_enrolled, created_at, active FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
    subjects = connection.execute("SELECT * FROM subjects WHERE class_id = ? AND active = 1 ORDER BY subject_code", (class_id,)).fetchall()
    today = date.today().isoformat()

    present_today = connection.execute(
        "SELECT COUNT(*) AS total FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE s.class_id = ? AND a.date = ? AND a.status = 'Present'",
        (class_id, today),
    ).fetchone()["total"]
    absent_today = connection.execute(
        "SELECT COUNT(*) AS total FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE s.class_id = ? AND a.date = ? AND a.status = 'Absent'",
        (class_id, today),
    ).fetchone()["total"]
    manual_ocr_today = connection.execute(
        "SELECT COUNT(*) AS total FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE s.class_id = ? AND a.date = ? AND a.method IN ('MANUAL', 'SHEET')",
        (class_id, today),
    ).fetchone()["total"]

    low_attendance = 0
    for student in students:
        stats = student_summary(student["id"])
        if stats["total"] and stats["percentage"] < 75:
            low_attendance += 1

    notifications = connection.execute(
        "SELECT * FROM notifications WHERE class_id = ? ORDER BY created_at DESC LIMIT 5",
        (class_id,),
    ).fetchall()
    subject_analytics = connection.execute(
        "SELECT s.id, s.subject_code, s.subject_name, COUNT(a.id) AS marked, SUM(CASE WHEN a.status = 'Present' THEN 1 ELSE 0 END) AS present, ROUND(100.0 * SUM(CASE WHEN a.status = 'Present' THEN 1 ELSE 0 END) / NULLIF(COUNT(a.id), 0), 1) AS rate FROM subjects s LEFT JOIN attendance a ON a.subject_id = s.id WHERE s.class_id = ? AND s.active = 1 GROUP BY s.id ORDER BY s.subject_code",
        (class_id,),
    ).fetchall()
    daily_analytics = connection.execute(
        "SELECT date, COUNT(*) AS marked, SUM(CASE WHEN status = 'Present' THEN 1 ELSE 0 END) AS present FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE s.class_id = ? AND date >= date('now', '-6 days') GROUP BY date ORDER BY date",
        (class_id,),
    ).fetchall()
    attendance_rows = connection.execute(
        "SELECT a.id, a.date, a.status, a.method, a.confidence, s.subject_name, st.name, st.student_uid, st.roll_no FROM attendance a JOIN subjects s ON s.id = a.subject_id JOIN students st ON st.id = a.student_id WHERE s.class_id = ? ORDER BY a.date DESC, a.id DESC LIMIT 100",
        (class_id,),
    ).fetchall()
    method_analytics = connection.execute(
        "SELECT a.method, COUNT(*) AS total FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE s.class_id = ? GROUP BY a.method",
        (class_id,),
    ).fetchall()
    connection.close()
    return {
        "class": dict(class_row) if class_row else None,
        "stats": {
            "total_students": len(students),
            "total_subjects": len(subjects),
            "present_today": present_today,
            "absent_today": absent_today,
            "manual_ocr_today": manual_ocr_today,
            "low_attendance": low_attendance,
        },
        "students": [dict(row) for row in students],
        "subjects": [dict(row) for row in subjects],
        "notifications": [dict(row) for row in notifications],
        "analytics": {
            "by_subject": [dict(row) for row in subject_analytics],
            "by_day": [dict(row) for row in daily_analytics],
            "methods": {row["method"]: row["total"] for row in method_analytics},
        },
        "recent_attendance": [dict(row) for row in attendance_rows],
    }


def run_ocr_on_image(image):
    return pytesseract.image_to_string(image).strip()


def match_reference_face(image, allowed_uids=None):
    face_crop, error = detect_face_crop(image)
    if error:
        return {"recognized": False, "reason": error}

    face_dir = Path("face_data")
    student_uids = sorted({path.stem.split("__", 1)[0] for path in face_dir.glob("*__*.jpg")})
    if allowed_uids is not None:
        student_uids = [student_uid for student_uid in student_uids if student_uid in allowed_uids]
    samples = []
    labels = []
    label_to_uid = {}
    for label, student_uid in enumerate(student_uids):
        label_to_uid[label] = student_uid
        for file_path in face_dir.glob(f"{student_uid}__*.jpg"):
            template = cv2.imread(str(file_path), cv2.IMREAD_GRAYSCALE)
            if template is not None:
                samples.append(cv2.resize(template, (200, 200)))
                labels.append(label)

    if not samples:
        return {"recognized": False, "reason": "No students have enrolled face samples yet."}

    recognizer = cv2.face.LBPHFaceRecognizer_create()
    recognizer.train(samples, np.asarray(labels, dtype=np.int32))
    label, distance = recognizer.predict(face_crop)
    confidence = max(0.0, min(1.0, 1.0 - float(distance) / 100.0))
    if distance > 65:
        return {"recognized": False, "reason": "Face did not match an enrolled student.", "confidence": round(confidence, 3)}
    return {"recognized": True, "student_uid": label_to_uid[label], "confidence": round(confidence, 3)}


@app.route("/api/faces", methods=["GET", "POST"])
def api_faces():
    if session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can manage face enrollment."}), 403
    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Select an assigned class first."}), 409

    face_dir = Path("face_data")
    face_dir.mkdir(exist_ok=True)
    connection = get_db_connection()
    if request.method == "GET":
        students = connection.execute("SELECT id, student_uid, name, roll_no, face_enrolled FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
        connection.close()
        face_list = []
        for student in students:
            count = len(list(face_dir.glob(f"{student['student_uid']}__*.jpg")))
            face_list.append({**dict(student), "sample_count": count, "face_enrolled": count > 0})
        return jsonify({"success": True, "students": face_list})

    student_id = request.form.get("student_id", type=int)
    image_file = request.files.get("image")
    if not student_id or not image_file:
        connection.close()
        return jsonify({"success": False, "message": "Choose a student and capture an enrollment image."}), 400
    student = connection.execute("SELECT id, student_uid FROM students WHERE id = ? AND class_id = ? AND active = 1", (student_id, class_id)).fetchone()
    if student is None:
        connection.close()
        return jsonify({"success": False, "message": "Student is not in the selected class."}), 404
    image_bytes = image_file.read()
    image = cv2.imdecode(np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        connection.close()
        return jsonify({"success": False, "message": "Could not read the camera image."}), 400
    face_crop, error = detect_face_crop(image)
    if error:
        connection.close()
        return jsonify({"success": False, "message": error}), 400

    sample_path = face_dir / f"{student['student_uid']}__{uuid.uuid4().hex}.jpg"
    if not cv2.imwrite(str(sample_path), face_crop):
        connection.close()
        return jsonify({"success": False, "message": "Could not store the face sample."}), 500
    connection.execute("UPDATE students SET face_enrolled = 1, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (student_id,))
    connection.execute("INSERT INTO audit_logs (user_id, action, entity_type, entity_id, new_value, ip_address) VALUES (?, 'FACE_ENROLLED', 'student', ?, ?, ?)", (session.get("user_id"), student_id, sample_path.name, request.remote_addr))
    connection.commit()
    connection.close()
    return jsonify({"success": True, "student_id": student_id, "sample_count": len(list(face_dir.glob(f"{student['student_uid']}__*.jpg")))})


@app.route("/api/faces/<int:student_id>", methods=["DELETE"])
def api_delete_faces(student_id):
    if session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can remove face enrollment."}), 403
    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Select an assigned class first."}), 409
    connection = get_db_connection()
    student = connection.execute("SELECT student_uid FROM students WHERE id = ? AND class_id = ?", (student_id, class_id)).fetchone()
    if student is None:
        connection.close()
        return jsonify({"success": False, "message": "Student is not in the selected class."}), 404
    for path in Path("face_data").glob(f"{student['student_uid']}__*.jpg"):
        path.unlink(missing_ok=True)
    connection.execute("UPDATE students SET face_enrolled = 0, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (student_id,))
    connection.execute("INSERT INTO audit_logs (user_id, action, entity_type, entity_id, ip_address) VALUES (?, 'FACE_ENROLLMENT_REMOVED', 'student', ?, ?)", (session.get("user_id"), student_id, request.remote_addr))
    connection.commit()
    connection.close()
    return jsonify({"success": True})


def normalize_ocr_text(value):
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def match_roster_student(text, students):
    tokens = re.findall(r"[A-Z0-9]+", str(text or "").upper())
    normalized_text = normalize_ocr_text(text)
    best_student = None
    best_score = 0
    for student in students:
        aliases = [student["student_uid"], student["roll_no"], student["name"]]
        for alias in filter(None, aliases):
            normalized_alias = normalize_ocr_text(alias)
            if not normalized_alias:
                continue
            if normalized_alias in normalized_text:
                score = 100
            else:
                alias_tokens = re.findall(r"[A-Z0-9]+", str(alias).upper())
                candidates = tokens
                if len(alias_tokens) > 1:
                    candidates = [
                        " ".join(tokens[start : start + size])
                        for size in range(max(1, len(alias_tokens) - 1), len(alias_tokens) + 2)
                        for start in range(max(0, len(tokens) - size + 1))
                    ]
                score = max(
                    (difflib.SequenceMatcher(None, normalize_ocr_text(candidate), normalized_alias).ratio() * 100 for candidate in candidates),
                    default=0,
                )
            if score > best_score:
                best_student, best_score = student, round(score)
    return best_student, best_score


def parse_attendance_status(text):
    value = str(text or "").upper()
    statuses = set(re.findall(r"\b(PRESENT|ABSENT|LATE)\b", value))
    if len(statuses) == 1:
        return statuses.pop().title(), 100
    if re.search(r"\b(TRUE|YES|SIGNED|CHECKED)\b|[✓✔☑]", value):
        return "Present", 95
    if re.search(r"\b(FALSE|NO|UNSIGNED|UNCHECKED)\b", value):
        return "Absent", 95
    return "Review", 0


def detect_checkmark_statuses(image, expected_rows):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    binary = cv2.threshold(gray, 180, 255, cv2.THRESH_BINARY_INV)[1]
    contours, _ = cv2.findContours(binary, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    boxes = []
    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        if 15 <= width <= 100 and 15 <= height <= 100 and 0.75 <= width / height <= 1.3:
            center = (x + width // 2, y + height // 2)
            if not any(abs(center[0] - old[0]) <= 2 and abs(center[1] - old[1]) <= 2 for old in boxes):
                boxes.append(center)

    y_tolerance = max(4, image.shape[0] // 150)
    row_groups = []
    for x, y in sorted(boxes, key=lambda point: (point[1], point[0])):
        if not row_groups or abs(y - row_groups[-1]["center"]) > y_tolerance:
            row_groups.append({"center": y, "boxes": [(x, y)]})
        else:
            group = row_groups[-1]
            group["boxes"].append((x, y))
            group["center"] = round(sum(point[1] for point in group["boxes"]) / len(group["boxes"]))

    if len(row_groups) < expected_rows:
        return {}

    detected = {}
    for row_index, group in enumerate(row_groups[:expected_rows]):
        row_boxes = sorted(group["boxes"])
        if len(row_boxes) < 2:
            continue
        present_box, absent_box = row_boxes[0], row_boxes[-1]
        checked = []
        for x_center, y_center in (present_box, absent_box):
            radius = max(3, image.shape[1] // 100)
            x1 = max(0, x_center - radius)
            x2 = min(image.shape[1], x_center + radius + 1)
            y1 = max(0, y_center - radius)
            y2 = min(image.shape[0], y_center + radius + 1)
            checked.append(cv2.countNonZero(binary[y1:y2, x1:x2]) >= 3)
        if checked[0] != checked[1]:
            detected[row_index] = "Present" if checked[0] else "Absent"
    return detected


def build_sheet_rows(lines, students):
    results = []
    assigned_student_ids = set()
    for line in lines:
        student, identity_confidence = match_roster_student(line, students)
        if student is None or identity_confidence < 70 or student["id"] in assigned_student_ids:
            continue
        status, status_confidence = parse_attendance_status(line)
        confidence = min(identity_confidence, status_confidence)
        results.append({
            "student_id": student["id"],
            "student_uid": student["student_uid"],
            "roll_no": student["roll_no"],
            "name": student["name"],
            "matched_text": str(line)[:160],
            "status": status if confidence >= 85 else "Review",
            "confidence": confidence,
            "signature_present": status == "Present" if status != "Review" else None,
            "detected_by_vision": False,
            "needs_verification": confidence < 85,
            "verified": False,
        })
        assigned_student_ids.add(student["id"])

    for student in students:
        if student["id"] in assigned_student_ids:
            continue
        results.append({
            "student_id": student["id"],
            "student_uid": student["student_uid"],
            "roll_no": student["roll_no"],
            "name": student["name"],
            "matched_text": "",
            "status": "Review",
            "confidence": 0,
            "signature_present": None,
            "detected_by_vision": False,
            "needs_verification": True,
            "verified": False,
        })
    return results


def read_tabular_sheet(file_path, extension):
    if extension == "csv":
        with open(file_path, newline="", encoding="utf-8-sig") as sheet_file:
            return list(csv.reader(sheet_file))
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def extract_tabular_sheet_results(file_path, extension, class_id, subject_id, attendance_date):
    rows = read_tabular_sheet(file_path, extension)
    if len(rows) < 2:
        return {"status": "error", "message": "The spreadsheet must include a header row and at least one student row."}

    headers = [normalize_ocr_text(value) for value in rows[0]]
    identity_columns = {"ROLLNO", "ROLLNUMBER", "STUDENTID", "STUDENTUID", "UID", "NAME", "STUDENTNAME"}
    status_columns = {"STATUS", "ATTENDANCE", "PRESENT", "SIGNATURE"}
    identity_index = next((index for index, header in enumerate(headers) if header in identity_columns), None)
    status_index = next((index for index, header in enumerate(headers) if header in status_columns), None)
    if identity_index is None or status_index is None:
        return {"status": "error", "message": "Include a student name, roll number, or student ID column and an attendance/status column."}
    signature_column = headers[status_index] == "SIGNATURE"

    connection = get_db_connection()
    students = connection.execute(
        "SELECT id, student_uid, roll_no, name FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no",
        (class_id,),
    ).fetchall()
    connection.close()
    parsed_lines = []
    for row in rows[1:]:
        identity = str(row[identity_index] or "").strip() if identity_index < len(row) else ""
        raw_status = str(row[status_index] or "").strip() if status_index < len(row) else ""
        if identity:
            if signature_column:
                raw_status = "Present" if raw_status else "Absent"
            parsed_lines.append(f"{identity} {raw_status}")
    results = build_sheet_rows(parsed_lines, students)
    return {
        "status": "ok",
        "ocr_text": "",
        "results": results,
        "detected": sum(row["status"] != "Review" for row in results),
        "subject_id": subject_id,
        "attendance_date": attendance_date,
    }


def extract_sheet_results(image_path, class_id, subject_id, attendance_date):
    image = cv2.imread(image_path)
    if image is None:
        return {"status": "error", "message": "Could not read the uploaded sheet image."}

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    threshold = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    points = cv2.findNonZero(threshold)
    if points is not None:
        angle = cv2.minAreaRect(points)[-1]
        angle = -(90 + angle) if angle < -45 else -angle
        if 0.5 < abs(angle) <= 15:
            height, width = image.shape[:2]
            matrix = cv2.getRotationMatrix2D((width / 2, height / 2), angle, 1.0)
            image = cv2.warpAffine(image, matrix, (width, height), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)
    text = run_ocr_on_image(image)
    connection = get_db_connection()
    students = connection.execute(
        "SELECT id, student_uid, roll_no, name FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no",
        (class_id,),
    ).fetchall()
    connection.close()

    results = build_sheet_rows(text.splitlines(), students)
    marks = detect_checkmark_statuses(image, len(students))
    row_by_student = {row["student_id"]: row for row in results}
    for index, student in enumerate(students):
        mark_status = marks.get(index)
        if not mark_status:
            continue
        row = row_by_student[student["id"]]
        if row["status"] not in {"Review", mark_status}:
            row["status"] = "Review"
            row["confidence"] = 50
            row["needs_verification"] = True
            row["signature_present"] = None
            continue
        row["status"] = mark_status
        row["confidence"] = 95
        row["needs_verification"] = False
        row["signature_present"] = mark_status == "Present"
        row["detected_by_vision"] = True
    recognized = sum(row["status"] != "Review" for row in results)

    return {
        "status": "ok",
        "ocr_text": text[:2000],
        "results": results,
        "detected": recognized,
        "subject_id": subject_id,
        "attendance_date": attendance_date,
    }


UPLOAD_FOLDER = os.environ.get("UPLOAD_FOLDER", "uploads")
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg"}
SHEET_EXTENSIONS = {"csv", "xlsx"}
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("MAX_UPLOAD_SIZE", 10 * 1024 * 1024))
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

create_tables()
seed_demo_data()


# ==================================================
# HELPERS
# ==================================================

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in IMAGE_EXTENSIONS


def get_selected_class():
    class_id = session.get("selected_class_id")
    if not class_id:
        return None
    row = get_db_connection().execute(
        """
        SELECT c.id, c.year, c.department, c.division, a.academic_year
        FROM classes c
        JOIN academic_years a ON a.id = c.academic_year_id
        WHERE c.id = ?
        """,
        (class_id,),
    ).fetchone()
    return row


def set_user_session(user_row, role):
    session["user_id"] = user_row["id"]
    session["email"] = user_row["email"]
    session["role"] = role
    session["name"] = user_row["name"]


def generate_student_uid(class_id, name=""):
    connection = get_db_connection()
    class_row = connection.execute(
        "SELECT c.department, c.year, a.academic_year FROM classes c JOIN academic_years a ON a.id = c.academic_year_id WHERE c.id = ?",
        (class_id,),
    ).fetchone()
    dept_prefix = "".join(part[0].upper() for part in class_row["department"].split()[:2])
    academic_year_start = class_row["academic_year"].split("-")[0]
    next_no = connection.execute(
        "SELECT COUNT(*) + 1 AS next_no FROM students WHERE class_id = ?",
        (class_id,),
    ).fetchone()["next_no"]
    student_uid = f"{dept_prefix}{academic_year_start}-{next_no:03d}"
    while connection.execute("SELECT 1 FROM students WHERE student_uid = ?", (student_uid,)).fetchone():
        next_no += 1
        student_uid = f"{dept_prefix}{academic_year_start}-{next_no:03d}"
    connection.close()
    return student_uid


def create_notification(student_id, class_id, subject_id, title, message, notification_type="ATTENDANCE", priority="NORMAL", connection=None):
    owns_connection = connection is None
    if owns_connection:
        connection = get_db_connection()
    connection.execute(
        """
        INSERT INTO notifications (student_id, class_id, subject_id, title, message, type, priority, is_read)
        VALUES (?, ?, ?, ?, ?, ?, ?, 0)
        """,
        (student_id, class_id, subject_id, title, message, notification_type, priority),
    )
    if owns_connection:
        connection.commit()
        connection.close()


def create_attendance_notification(connection, student_id, subject_name, attendance_date, status):
    student = connection.execute("SELECT id, class_id FROM students WHERE id = ?", (student_id,)).fetchone()
    if student is None:
        return
    subject = connection.execute("SELECT id FROM subjects WHERE class_id = ? AND subject_name = ? LIMIT 1", (student["class_id"], subject_name)).fetchone()
    create_notification(
        student_id,
        student["class_id"],
        subject["id"] if subject else None,
        "Attendance Marked",
        f"Your {subject_name} attendance for {attendance_date} has been marked {status}.",
        "ATTENDANCE",
        "NORMAL",
        connection=connection,
    )


def create_low_attendance_notification(connection, student_id):
    student = connection.execute("SELECT id, class_id FROM students WHERE id = ?", (student_id,)).fetchone()
    if student is None:
        return
    total = connection.execute("SELECT COUNT(*) AS total FROM attendance WHERE student_id = ?", (student_id,)).fetchone()["total"]
    present = connection.execute("SELECT COUNT(*) AS total FROM attendance WHERE student_id = ? AND status = 'Present'", (student_id,)).fetchone()["total"]
    if total == 0:
        return
    percentage = (present / total) * 100
    if percentage < 75:
        create_notification(
            student_id,
            student["class_id"],
            None,
            "Low Attendance Alert",
            f"Your current attendance is {percentage:.2f}%. It is below the required 75% threshold.",
            "LOW_ATTENDANCE",
            "HIGH",
            connection=connection,
        )


def student_summary(student_id):
    connection = get_db_connection()
    total = connection.execute("SELECT COUNT(*) AS total FROM attendance WHERE student_id = ?", (student_id,)).fetchone()["total"]
    present = connection.execute("SELECT COUNT(*) AS total FROM attendance WHERE student_id = ? AND status = 'Present'", (student_id,)).fetchone()["total"]
    absent = connection.execute("SELECT COUNT(*) AS total FROM attendance WHERE student_id = ? AND status = 'Absent'", (student_id,)).fetchone()["total"]
    late = connection.execute("SELECT COUNT(*) AS total FROM attendance WHERE student_id = ? AND status = 'Late'", (student_id,)).fetchone()["total"]
    percentage = round((present / total) * 100, 2) if total else 0
    connection.close()
    return {"total": total, "present": present, "absent": absent, "late": late, "percentage": percentage}


@app.route("/api/login", methods=["POST"])
def api_login():
    payload = request.get_json(silent=True) or request.form
    identifier = str(payload.get("identifier", "")).strip()
    password = str(payload.get("password", ""))
    role = str(payload.get("role", "")).lower()

    if not identifier or not password or not role:
        return jsonify({"success": False, "message": "Invalid credentials."}), 400

    connection = get_db_connection()
    user = None
    if role == "student":
        user = connection.execute(
            "SELECT s.*, u.email, u.password_hash FROM students s LEFT JOIN users u ON u.email = s.email WHERE (s.student_uid = ? OR s.email = ?) AND s.active = 1 LIMIT 1",
            (identifier, identifier),
        ).fetchone()
        if user and user["password_hash"] and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            session["role"] = "student"
            session["email"] = user["email"]
            session["name"] = user["name"]
            session["student_id"] = user["id"]
            connection.close()
            return jsonify({"success": True, "role": "student", "redirect": "/student-dashboard"})
    else:
        user = connection.execute(
            "SELECT * FROM users WHERE email = ? AND role = ? LIMIT 1",
            (identifier, role),
        ).fetchone()
        if user and user["password_hash"] and check_password_hash(user["password_hash"], password):
            set_user_session(user, role)
            connection.close()
            return jsonify({"success": True, "role": role, "redirect": "/class-selection"})

    connection.close()
    return jsonify({"success": False, "message": "Invalid email or password."}), 401


@app.route("/api/dashboard")
def api_dashboard():
    role = session.get("role")
    if role == "student":
        connection = get_db_connection()
        student = connection.execute("SELECT id, student_uid, roll_no, name, email, phone, class_id, face_enrolled, created_at, active FROM students WHERE id = ?", (session.get("student_id"),)).fetchone()
        if student is None:
            connection.close()
            return jsonify({"success": False, "message": "Student not found."}), 404
        class_row = connection.execute(
            "SELECT c.id, c.year, c.department, c.division, a.academic_year FROM classes c JOIN academic_years a ON a.id = c.academic_year_id WHERE c.id = ?",
            (student["class_id"],),
        ).fetchone()
        summary = student_summary(student["id"])
        subject_rows = connection.execute(
            "SELECT s.id, s.subject_code, s.subject_name, COUNT(a.id) AS total, SUM(CASE WHEN a.status = 'Present' THEN 1 ELSE 0 END) AS present, SUM(CASE WHEN a.status = 'Absent' THEN 1 ELSE 0 END) AS absent, SUM(CASE WHEN a.status = 'Late' THEN 1 ELSE 0 END) AS late FROM subjects s LEFT JOIN attendance a ON a.subject_id = s.id AND a.student_id = ? WHERE s.class_id = ? AND s.active = 1 GROUP BY s.id ORDER BY s.subject_code",
            (student["id"], student["class_id"]),
        ).fetchall()
        subject_cards = []
        for row in subject_rows:
            record = dict(row)
            record["total"] = record["total"] or 0
            record["present"] = record["present"] or 0
            record["absent"] = record["absent"] or 0
            record["late"] = record["late"] or 0
            record["percentage"] = round(100 * record["present"] / record["total"], 1) if record["total"] else 0
            subject_cards.append(record)
        recent = connection.execute(
            "SELECT a.date, a.status, a.method, a.confidence, s.subject_code, s.subject_name FROM attendance a JOIN subjects s ON s.id = a.subject_id WHERE a.student_id = ? ORDER BY a.date DESC, a.id DESC LIMIT 15",
            (student["id"],),
        ).fetchall()
        notifications = connection.execute(
            "SELECT id, title, message, type, priority, is_read, created_at FROM notifications WHERE student_id = ? ORDER BY created_at DESC LIMIT 6",
            (student["id"],),
        ).fetchall()
        connection.close()
        return jsonify({"success": True, "role": "student", "student": dict(student), "class": dict(class_row) if class_row else None, "summary": summary, "subjects": subject_cards, "recent_attendance": [dict(row) for row in recent], "notifications": [dict(row) for row in notifications]})

    if role in {"admin", "faculty"}:
        if not session.get("selected_class_id") or not class_access_allowed(session["selected_class_id"]):
            return jsonify({"success": False, "message": "Select a class assigned to your account."}), 409
        payload = build_dashboard_payload()
        return jsonify({"success": True, "role": role, "dashboard": payload})

    return jsonify({"success": False, "message": "Unauthorized."}), 401


@app.route("/api/class-selection", methods=["GET", "POST"])
def api_class_selection():
    if session.get("role") not in {"admin", "faculty"}:
        return jsonify({"success": False, "message": "Unauthorized."}), 401

    connection = get_db_connection()
    if request.method == "POST":
        payload = request.get_json(silent=True) or request.form
        academic_year_id = str(payload.get("academic_year_id", "")).strip()
        year = str(payload.get("year", "")).strip()
        department = str(payload.get("department", "")).strip()
        division = str(payload.get("division", "")).strip()
        if not all([academic_year_id, year, department, division]):
            return jsonify({"success": False, "message": "Please complete all class details."}), 400

        class_row = connection.execute(
            "SELECT id FROM classes WHERE academic_year_id = ? AND year = ? AND department = ? AND division = ? AND active = 1 LIMIT 1",
            (academic_year_id, year, department, division),
        ).fetchone()
        if class_row is None:
            if session.get("role") != "faculty":
                connection.close()
                return jsonify({"success": False, "message": "Only faculty can create new class divisions."}), 403
            year_exists = connection.execute("SELECT 1 FROM academic_years WHERE id = ?", (academic_year_id,)).fetchone()
            if year_exists is None:
                connection.close()
                return jsonify({"success": False, "message": "Choose a valid academic year."}), 400
            connection.execute(
                "INSERT INTO classes (academic_year_id, year, department, division) VALUES (?, ?, ?, ?)",
                (academic_year_id, year, department, division),
            )
            class_row = connection.execute(
                "SELECT id FROM classes WHERE academic_year_id = ? AND year = ? AND department = ? AND division = ? AND active = 1 LIMIT 1",
                (academic_year_id, year, department, division),
            ).fetchone()
            connection.commit()
        if session.get("role") == "faculty":
            faculty = connection.execute("SELECT id FROM faculty WHERE user_id = ?", (session.get("user_id"),)).fetchone()
            if faculty is None:
                connection.execute("INSERT INTO faculty (user_id, faculty_code) VALUES (?, ?)", (session.get("user_id"), f"FAC-{session.get('user_id')}"))
                faculty = connection.execute("SELECT id FROM faculty WHERE user_id = ?", (session.get("user_id"),)).fetchone()
            connection.execute("INSERT INTO faculty_class_assignments (faculty_id, class_id) SELECT ?, ? WHERE NOT EXISTS (SELECT 1 FROM faculty_class_assignments WHERE faculty_id = ? AND class_id = ?)", (faculty["id"], class_row["id"], faculty["id"], class_row["id"]))
            connection.commit()
        if not class_access_allowed(class_row["id"]):
            connection.close()
            return jsonify({"success": False, "message": "This class is not assigned to your account."}), 403
        session["selected_class_id"] = class_row["id"]
        connection.close()
        return jsonify({"success": True, "class_id": class_row["id"], "redirect": "/admin-dashboard"})

    academic_years = connection.execute("SELECT * FROM academic_years ORDER BY academic_year DESC").fetchall()
    classes = connection.execute("SELECT c.id, c.year, c.department, c.division, a.academic_year, c.academic_year_id FROM classes c JOIN academic_years a ON a.id = c.academic_year_id WHERE c.active = 1 ORDER BY a.academic_year DESC, c.department, c.year, c.division").fetchall()
    if session.get("role") == "faculty":
        archived_classes = connection.execute(
            "SELECT c.id, c.year, c.department, c.division, a.academic_year, c.academic_year_id FROM classes c JOIN academic_years a ON a.id = c.academic_year_id JOIN faculty_class_assignments fca ON fca.class_id = c.id JOIN faculty f ON f.id = fca.faculty_id WHERE c.active = 0 AND f.user_id = ? ORDER BY a.academic_year DESC, c.department, c.year, c.division",
            (session.get("user_id"),),
        ).fetchall()
    else:
        archived_classes = []
    connection.close()
    return jsonify({"success": True, "academic_years": [dict(row) for row in academic_years], "classes": [dict(row) for row in classes], "archived_classes": [dict(row) for row in archived_classes], "selected_class_id": session.get("selected_class_id")})


@app.route("/api/classes/<int:class_id>", methods=["PUT", "DELETE"])
def api_class_detail(class_id):
    if session.get("role") != "faculty" or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Only faculty can edit or archive classes."}), 403

    connection = get_db_connection()
    class_row = connection.execute("SELECT id FROM classes WHERE id = ? AND active = 1", (class_id,)).fetchone()
    if class_row is None:
        connection.close()
        return jsonify({"success": False, "message": "Class not found."}), 404

    if request.method == "DELETE":
        connection.execute("UPDATE classes SET active = 0 WHERE id = ?", (class_id,))
        connection.execute("INSERT INTO audit_logs (user_id, action, entity_type, entity_id, new_value, ip_address) VALUES (?, 'CLASS_ARCHIVED', 'class', ?, 'active=0', ?)", (session.get("user_id"), class_id, request.remote_addr))
        connection.commit()
        connection.close()
        if session.get("selected_class_id") == class_id:
            session.pop("selected_class_id", None)
        return jsonify({"success": True})

    payload = request.get_json(silent=True) or request.form
    academic_year_id = str(payload.get("academic_year_id", "")).strip()
    year = str(payload.get("year", "")).strip()
    department = str(payload.get("department", "")).strip()
    division = str(payload.get("division", "")).strip()
    if not all([academic_year_id, year, department, division]):
        connection.close()
        return jsonify({"success": False, "message": "Academic year, study year, department, and division are required."}), 400
    academic_year = connection.execute("SELECT 1 FROM academic_years WHERE id = ?", (academic_year_id,)).fetchone()
    if academic_year is None:
        connection.close()
        return jsonify({"success": False, "message": "Choose a valid academic year."}), 400
    duplicate = connection.execute("SELECT 1 FROM classes WHERE academic_year_id = ? AND year = ? AND department = ? AND division = ? AND id != ? AND active = 1", (academic_year_id, year, department, division, class_id)).fetchone()
    if duplicate:
        connection.close()
        return jsonify({"success": False, "message": "That academic year, department, study year, and division already exist."}), 409
    connection.execute("UPDATE classes SET academic_year_id = ?, year = ?, department = ?, division = ? WHERE id = ?", (academic_year_id, year, department, division, class_id))
    connection.execute("INSERT INTO audit_logs (user_id, action, entity_type, entity_id, new_value, ip_address) VALUES (?, 'CLASS_UPDATED', 'class', ?, ?, ?)", (session.get("user_id"), class_id, f"{academic_year_id}|{year}|{department}|{division}", request.remote_addr))
    connection.commit()
    connection.close()
    return jsonify({"success": True})


@app.route("/api/classes/<int:class_id>/restore", methods=["POST"])
def api_restore_class(class_id):
    if session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can restore classes."}), 403

    connection = get_db_connection()
    class_row = connection.execute(
        "SELECT c.id, c.academic_year_id, c.year, c.department, c.division FROM classes c JOIN faculty_class_assignments fca ON fca.class_id = c.id JOIN faculty f ON f.id = fca.faculty_id WHERE c.id = ? AND c.active = 0 AND f.user_id = ?",
        (class_id, session.get("user_id")),
    ).fetchone()
    if class_row is None:
        connection.close()
        return jsonify({"success": False, "message": "Archived class not found in your assignments."}), 404

    duplicate = connection.execute(
        "SELECT 1 FROM classes WHERE academic_year_id = ? AND year = ? AND department = ? AND division = ? AND active = 1",
        (class_row["academic_year_id"], class_row["year"], class_row["department"], class_row["division"]),
    ).fetchone()
    if duplicate:
        connection.close()
        return jsonify({"success": False, "message": "An active class with these details already exists."}), 409

    connection.execute("UPDATE classes SET active = 1 WHERE id = ?", (class_id,))
    connection.execute(
        "INSERT INTO audit_logs (user_id, action, entity_type, entity_id, new_value, ip_address) VALUES (?, 'CLASS_RESTORED', 'class', ?, 'active=1', ?)",
        (session.get("user_id"), class_id, request.remote_addr),
    )
    connection.commit()
    connection.close()
    return jsonify({"success": True})


@app.route("/api/students", methods=["GET", "POST"])
def api_students():
    if session.get("role") not in {"admin", "faculty"}:
        return jsonify({"success": False, "message": "Unauthorized."}), 401
    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": True, "students": []})
    if request.method == "POST" and session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can add students."}), 403
    connection = get_db_connection()
    if request.method == "POST":
        payload = request.get_json(silent=True) or request.form
        name = str(payload.get("name", "")).strip()
        roll_no = str(payload.get("roll_no", "")).strip()
        email = str(payload.get("email", "")).strip().lower()
        phone = str(payload.get("phone", "")).strip()
        if not all([name, roll_no, email]):
            connection.close()
            return jsonify({"success": False, "message": "Name, roll number, and email are required."}), 400
        duplicate_roll = connection.execute("SELECT 1 FROM students WHERE class_id = ? AND roll_no = ? AND active = 1", (class_id, roll_no)).fetchone()
        if duplicate_roll:
            connection.close()
            return jsonify({"success": False, "message": "That roll number is already used in this class."}), 409
        student_uid = generate_student_uid(class_id)
        try:
            password_hash = generate_password_hash(str(payload.get("password") or "student123"))
            connection.execute("INSERT INTO students (student_uid, roll_no, name, email, phone, class_id, password_hash, active) VALUES (?, ?, ?, ?, ?, ?, ?, 1)", (student_uid, roll_no, name, email, phone, class_id, password_hash))
            connection.execute("INSERT INTO users (name, email, password_hash, role) VALUES (?, ?, ?, 'student')", (name, email, password_hash))
            student_id = connection.execute("SELECT id FROM students WHERE student_uid = ?", (student_uid,)).fetchone()["id"]
            connection.commit()
        except sqlite3.IntegrityError:
            connection.rollback()
            connection.close()
            return jsonify({"success": False, "message": "That email or roll number is already in use."}), 409
        connection.close()
        return jsonify({"success": True, "student_id": student_id, "student_uid": student_uid}), 201
    include_inactive = request.args.get("include_inactive") == "true"
    active_filter = "" if include_inactive else "AND active = 1"
    rows = connection.execute(f"SELECT id, student_uid, roll_no, name, email, phone, class_id, face_enrolled, created_at, active FROM students WHERE class_id = ? {active_filter} ORDER BY active DESC, roll_no", (class_id,)).fetchall()
    connection.close()
    return jsonify({"success": True, "students": [dict(row) for row in rows]})


@app.route("/api/students/<int:student_id>", methods=["PUT", "DELETE"])
def api_student_detail(student_id):
    class_id = session.get("selected_class_id")
    if session.get("role") != "faculty" or not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Only faculty can edit or remove students."}), 403
    connection = get_db_connection()
    student = connection.execute("SELECT * FROM students WHERE id = ? AND class_id = ? AND active = 1", (student_id, class_id)).fetchone()
    if student is None:
        connection.close()
        return jsonify({"success": False, "message": "Student not found in this class."}), 404
    if request.method == "DELETE":
        connection.execute("UPDATE students SET active = 0, face_enrolled = 0, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (student_id,))
        connection.execute("UPDATE users SET updated_at = CURRENT_TIMESTAMP WHERE email = ? AND role = 'student'", (student["email"],))
        connection.execute("INSERT INTO audit_logs (user_id, action, entity_type, entity_id, old_value, ip_address) VALUES (?, 'STUDENT_DEACTIVATED', 'student', ?, ?, ?)", (session.get("user_id"), student_id, student["student_uid"], request.remote_addr))
        connection.commit()
        connection.close()
        for sample in Path("face_data").glob(f"{student['student_uid']}__*.jpg"):
            sample.unlink(missing_ok=True)
        return jsonify({"success": True})

    payload = request.get_json(silent=True) or request.form
    name = str(payload.get("name", "")).strip()
    roll_no = str(payload.get("roll_no", "")).strip()
    email = str(payload.get("email", "")).strip().lower()
    phone = str(payload.get("phone", "")).strip()
    if not all([name, roll_no, email]):
        connection.close()
        return jsonify({"success": False, "message": "Name, roll number, and email are required."}), 400
    duplicate_roll = connection.execute("SELECT 1 FROM students WHERE class_id = ? AND roll_no = ? AND id != ? AND active = 1", (class_id, roll_no, student_id)).fetchone()
    if duplicate_roll:
        connection.close()
        return jsonify({"success": False, "message": "That roll number is already used in this class."}), 409
    try:
        connection.execute("UPDATE students SET name = ?, roll_no = ?, email = ?, phone = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (name, roll_no, email, phone, student_id))
        connection.execute("UPDATE users SET name = ?, email = ?, updated_at = CURRENT_TIMESTAMP WHERE email = ? AND role = 'student'", (name, email, student["email"]))
        connection.commit()
    except sqlite3.IntegrityError:
        connection.rollback()
        connection.close()
        return jsonify({"success": False, "message": "That email or roll number is already in use."}), 409
    connection.close()
    return jsonify({"success": True})


@app.route("/api/students/<int:student_id>/restore", methods=["POST"])
def api_restore_student(student_id):
    class_id = session.get("selected_class_id")
    if session.get("role") != "faculty" or not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Only faculty can restore students in the selected class."}), 403

    connection = get_db_connection()
    student = connection.execute(
        "SELECT id, roll_no FROM students WHERE id = ? AND class_id = ? AND active = 0",
        (student_id, class_id),
    ).fetchone()
    if student is None:
        connection.close()
        return jsonify({"success": False, "message": "Archived student not found in this class."}), 404

    duplicate_roll = connection.execute(
        "SELECT 1 FROM students WHERE class_id = ? AND roll_no = ? AND active = 1",
        (class_id, student["roll_no"]),
    ).fetchone()
    if duplicate_roll:
        connection.close()
        return jsonify({"success": False, "message": "An active student in this class already uses that roll number."}), 409

    connection.execute(
        "UPDATE students SET active = 1, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (student_id,),
    )
    connection.execute(
        "INSERT INTO audit_logs (user_id, action, entity_type, entity_id, new_value, ip_address) VALUES (?, 'STUDENT_RESTORED', 'student', ?, 'active=1', ?)",
        (session.get("user_id"), student_id, request.remote_addr),
    )
    connection.commit()
    connection.close()
    return jsonify({"success": True})


@app.route("/api/subjects", methods=["GET", "POST"])
def api_subjects():
    if session.get("role") not in {"admin", "faculty"}:
        return jsonify({"success": False, "message": "Unauthorized."}), 401
    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": True, "subjects": []})
    if request.method == "POST" and session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can add subjects."}), 403
    connection = get_db_connection()
    if request.method == "POST":
        payload = request.get_json(silent=True) or request.form
        subject_code = str(payload.get("subject_code", "")).strip().upper()
        subject_name = str(payload.get("subject_name", "")).strip()
        lecture_type = str(payload.get("lecture_type", "THEORY")).strip().upper()
        if not subject_code or not subject_name or lecture_type not in {"THEORY", "PRACTICAL", "LAB", "TUTORIAL"}:
            connection.close()
            return jsonify({"success": False, "message": "Provide a code, name, and valid subject type."}), 400
        try:
            total_hours = max(0, int(payload.get("total_hours", 0)))
        except (TypeError, ValueError):
            connection.close()
            return jsonify({"success": False, "message": "Total hours must be a number."}), 400
        faculty = connection.execute("SELECT id FROM faculty WHERE user_id = ?", (session.get("user_id"),)).fetchone() if session.get("role") == "faculty" else None
        try:
            cursor = connection.execute("INSERT INTO subjects (subject_code, subject_name, faculty_id, class_id, lecture_type, total_hours, active) VALUES (?, ?, ?, ?, ?, ?, 1)", (subject_code, subject_name, faculty["id"] if faculty else None, class_id, lecture_type, total_hours))
            connection.commit()
        except sqlite3.IntegrityError:
            connection.rollback()
            connection.close()
            return jsonify({"success": False, "message": "Could not add the subject. Check its code and try again."}), 409
        subject_id = cursor.lastrowid
        connection.close()
        return jsonify({"success": True, "subject_id": subject_id}), 201
    rows = connection.execute("SELECT s.*, f.faculty_code FROM subjects s LEFT JOIN faculty f ON f.id = s.faculty_id WHERE s.class_id = ? AND s.active = 1 ORDER BY s.subject_code", (class_id,)).fetchall()
    connection.close()
    return jsonify({"success": True, "subjects": [dict(row) for row in rows]})


@app.route("/api/subjects/<int:subject_id>", methods=["PUT", "DELETE"])
def api_subject_detail(subject_id):
    class_id = session.get("selected_class_id")
    if session.get("role") != "faculty" or not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Only faculty can edit or remove subjects."}), 403
    connection = get_db_connection()
    subject = connection.execute("SELECT * FROM subjects WHERE id = ? AND class_id = ? AND active = 1", (subject_id, class_id)).fetchone()
    if subject is None:
        connection.close()
        return jsonify({"success": False, "message": "Subject not found in this class."}), 404
    if request.method == "DELETE":
        connection.execute("UPDATE subjects SET active = 0 WHERE id = ?", (subject_id,))
        connection.commit()
        connection.close()
        return jsonify({"success": True})
    payload = request.get_json(silent=True) or request.form
    subject_code = str(payload.get("subject_code", "")).strip().upper()
    subject_name = str(payload.get("subject_name", "")).strip()
    lecture_type = str(payload.get("lecture_type", "THEORY")).strip().upper()
    if not subject_code or not subject_name or lecture_type not in {"THEORY", "PRACTICAL", "LAB", "TUTORIAL"}:
        connection.close()
        return jsonify({"success": False, "message": "Provide a code, name, and valid subject type."}), 400
    try:
        total_hours = max(0, int(payload.get("total_hours", 0)))
        connection.execute("UPDATE subjects SET subject_code = ?, subject_name = ?, lecture_type = ?, total_hours = ? WHERE id = ?", (subject_code, subject_name, lecture_type, total_hours, subject_id))
        connection.commit()
    except (sqlite3.IntegrityError, TypeError, ValueError):
        connection.rollback()
        connection.close()
        return jsonify({"success": False, "message": "Could not update the subject. Check the details."}), 409
    connection.close()
    return jsonify({"success": True})


@app.route("/api/attendance", methods=["GET", "POST"])
def api_attendance():
    class_id = session.get("selected_class_id")
    if session.get("role") != "faculty" or not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Faculty access to an active class is required."}), 403

    if request.method == "GET":
        subject_id = request.args.get("subject_id", type=int)
        attendance_date = request.args.get("date", date.today().isoformat())
        try:
            date.fromisoformat(attendance_date)
        except ValueError:
            return jsonify({"success": False, "message": "Date must use YYYY-MM-DD format."}), 400
        connection = get_db_connection()
        subjects = connection.execute("SELECT id, subject_code, subject_name FROM subjects WHERE class_id = ? AND active = 1 ORDER BY subject_code", (class_id,)).fetchall()
        if subject_id:
            subject = connection.execute("SELECT id FROM subjects WHERE id = ? AND class_id = ? AND active = 1", (subject_id, class_id)).fetchone()
            if subject is None:
                connection.close()
                return jsonify({"success": False, "message": "Subject is not in the selected class."}), 404
        students = connection.execute(
            "SELECT st.id, st.student_uid, st.roll_no, st.name, st.email, st.phone, a.status, a.method FROM students st LEFT JOIN attendance a ON a.student_id = st.id AND a.subject_id = ? AND a.date = ? WHERE st.class_id = ? AND st.active = 1 ORDER BY st.roll_no",
            (subject_id, attendance_date, class_id),
        ).fetchall()
        connection.close()
        return jsonify({"success": True, "date": attendance_date, "subject_id": subject_id, "subjects": [dict(row) for row in subjects], "students": [dict(row) for row in students]})

    payload = request.get_json(silent=True) or {}
    try:
        subject_id = int(payload.get("subject_id") or 0)
    except (TypeError, ValueError):
        subject_id = 0
    attendance_date = str(payload.get("date", "")).strip()
    records = payload.get("records", [])
    if not subject_id or not records:
        return jsonify({"success": False, "message": "Choose a subject and mark at least one student."}), 400
    try:
        date.fromisoformat(attendance_date)
    except ValueError:
        return jsonify({"success": False, "message": "Choose a valid attendance date."}), 400

    connection = get_db_connection()
    subject = connection.execute("SELECT id, subject_name FROM subjects WHERE id = ? AND class_id = ? AND active = 1", (subject_id, class_id)).fetchone()
    if subject is None:
        connection.close()
        return jsonify({"success": False, "message": "Subject is not in the selected class."}), 404
    saved = 0
    for record in records:
        student_id = record.get("student_id")
        status = record.get("status")
        if not student_id or status not in {"Present", "Absent", "Late"}:
            continue
        student = connection.execute("SELECT id FROM students WHERE id = ? AND class_id = ? AND active = 1", (student_id, class_id)).fetchone()
        if student is None:
            continue
        connection.execute(
            "INSERT INTO attendance (student_id, subject_id, date, status, method, confidence, recognized_by) VALUES (?, ?, ?, ?, 'MANUAL', 1.0, ?) ON CONFLICT(student_id, subject_id, date) DO UPDATE SET status = excluded.status, method = 'MANUAL', confidence = 1.0, recognized_by = excluded.recognized_by, updated_at = CURRENT_TIMESTAMP",
            (student_id, subject_id, attendance_date, status, session.get("name", "FACULTY")),
        )
        create_attendance_notification(connection, student_id, subject["subject_name"], attendance_date, status)
        create_low_attendance_notification(connection, student_id)
        saved += 1
    connection.commit()
    connection.close()
    return jsonify({"success": True, "saved": saved, "date": attendance_date})


@app.route("/api/camera-attendance", methods=["POST"])
def api_camera_attendance():
    if session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can mark camera attendance."}), 403

    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Select an assigned class first."}), 409
    subject_id = request.form.get("subject_id")
    attendance_date = request.form.get("attendance_date") or date.today().isoformat()
    image_file = request.files.get("image")
    if not image_file or not subject_id:
        return jsonify({"success": False, "message": "Subject and camera image are required."}), 400

    image = cv2.imdecode(np.frombuffer(image_file.read(), dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        return jsonify({"success": False, "message": "Could not read camera image."}), 400

    connection = get_db_connection()
    subject = connection.execute("SELECT id, subject_name FROM subjects WHERE id = ? AND class_id = ? AND active = 1", (subject_id, class_id)).fetchone()
    if subject is None:
        connection.close()
        return jsonify({"success": False, "message": "Subject is not in the selected class."}), 404
    allowed_uids = {row["student_uid"] for row in connection.execute("SELECT student_uid FROM students WHERE class_id = ? AND active = 1", (class_id,)).fetchall()}
    connection.close()

    result = match_reference_face(image, allowed_uids)
    if not result.get("recognized"):
        return jsonify({"success": False, "message": result.get("reason", "Face recognition failed."), "confidence": result.get("confidence", 0)})

    student_uid = result["student_uid"]
    connection = get_db_connection()
    student = connection.execute("SELECT * FROM students WHERE student_uid = ? AND class_id = ? AND active = 1 LIMIT 1", (student_uid, class_id)).fetchone()
    if student is None:
        connection.close()
        return jsonify({"success": False, "message": "Student template was not found."}), 404

    connection.execute(
        "INSERT INTO attendance (student_id, subject_id, date, status, method, confidence, recognized_by) VALUES (?, ?, ?, ?, 'CAMERA', ?, 'SYSTEM') ON CONFLICT(student_id, subject_id, date) DO UPDATE SET status = excluded.status, method = 'CAMERA', confidence = excluded.confidence, updated_at = CURRENT_TIMESTAMP",
        (student["id"], subject_id, attendance_date, "Present", result["confidence"]),
    )
    create_attendance_notification(connection, student["id"], subject["subject_name"], attendance_date, "Present")
    create_low_attendance_notification(connection, student["id"])
    connection.commit()
    connection.close()

    return jsonify({"success": True, "student": dict(student), "confidence": result["confidence"], "attendance_date": attendance_date})


@app.route("/api/attendance-sheet", methods=["POST"])
def api_attendance_sheet():
    if session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can upload attendance sheets."}), 403
    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Select an assigned class first."}), 409

    sheet_file = request.files.get("image") or request.files.get("file")
    if sheet_file is None or not sheet_file.filename:
        return jsonify({"success": False, "message": "An image, CSV, or Excel sheet is required."}), 400

    subject_id = request.form.get("subject_id")
    attendance_date = request.form.get("attendance_date") or date.today().isoformat()
    if not subject_id:
        return jsonify({"success": False, "message": "Subject is required."}), 400
    try:
        date.fromisoformat(attendance_date)
    except ValueError:
        return jsonify({"success": False, "message": "Attendance date must use YYYY-MM-DD format."}), 400
    extension = Path(sheet_file.filename).suffix.lower().lstrip(".")
    if extension not in IMAGE_EXTENSIONS | SHEET_EXTENSIONS:
        return jsonify({"success": False, "message": "Supported formats are JPG, PNG, CSV, and XLSX."}), 400

    connection = get_db_connection()
    subject = connection.execute("SELECT id FROM subjects WHERE id = ? AND class_id = ? AND active = 1", (subject_id, class_id)).fetchone()
    connection.close()
    if subject is None:
        return jsonify({"success": False, "message": "Subject is not in the selected class."}), 404
    filename = f"sheet_{uuid.uuid4()}.{extension}"
    file_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    sheet_file.save(file_path)
    try:
        if extension in SHEET_EXTENSIONS:
            result = extract_tabular_sheet_results(file_path, extension, class_id, subject_id, attendance_date)
        else:
            result = extract_sheet_results(file_path, class_id, subject_id, attendance_date)
    except pytesseract.TesseractNotFoundError:
        connection = get_db_connection()
        students = connection.execute("SELECT id, student_uid, roll_no, name FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
        connection.close()
        image = cv2.imread(file_path)
        marks = detect_checkmark_statuses(image, len(students)) if image is not None else {}
        fallback_rows = []
        for index, student in enumerate(students):
            status = marks.get(index, "Review")
            confidence = 95 if status != "Review" else 0
            fallback_rows.append({
                "student_id": student["id"],
                "student_uid": student["student_uid"],
                "roll_no": student["roll_no"],
                "name": student["name"],
                "matched_text": "",
                "status": status,
                "confidence": confidence,
                "signature_present": status == "Present" if status != "Review" else None,
                "detected_by_vision": status != "Review",
                "needs_verification": confidence < 85,
                "verified": False,
            })
        result = {
            "status": "review_required",
            "message": "OCR engine unavailable; the uploaded sheet is ready for manual review.",
            "ocr_text": "",
            "detected": 0,
            "subject_id": int(subject_id),
            "attendance_date": attendance_date,
            "results": fallback_rows,
        }
    except pytesseract.TesseractError as error:
        return jsonify({"success": False, "message": f"OCR could not process this image: {error}"}), 422
    except (OSError, UnicodeDecodeError, ValueError, csv.Error, InvalidFileException, zipfile.BadZipFile) as error:
        return jsonify({"success": False, "message": f"Could not read the uploaded sheet: {error}"}), 422
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)
    return jsonify(result)


@app.route("/api/attendance-sheet/confirm", methods=["POST"])
def api_confirm_sheet_attendance():
    if session.get("role") != "faculty":
        return jsonify({"success": False, "message": "Only faculty can confirm sheet attendance."}), 403
    class_id = session.get("selected_class_id")
    if not class_id or not class_access_allowed(class_id):
        return jsonify({"success": False, "message": "Select an assigned class first."}), 409
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"success": False, "message": "A valid JSON attendance payload is required."}), 400
    try:
        subject_id = int(payload.get("subject_id") or 0)
    except (TypeError, ValueError):
        subject_id = 0
    attendance_date = payload.get("attendance_date") or date.today().isoformat()
    records = payload.get("records", [])
    if not subject_id or not isinstance(records, list) or not records:
        return jsonify({"success": False, "message": "No attendance rows were provided."}), 400
    try:
        date.fromisoformat(attendance_date)
    except (TypeError, ValueError):
        return jsonify({"success": False, "message": "Attendance date must use YYYY-MM-DD format."}), 400

    student_ids = []
    for record in records:
        if not isinstance(record, dict):
            return jsonify({"success": False, "message": "Attendance rows must be valid objects."}), 400
        raw_student_id = record.get("student_id")
        if isinstance(raw_student_id, bool) or not isinstance(raw_student_id, (int, str)):
            student_id = 0
        else:
            try:
                student_id = int(raw_student_id)
            except ValueError:
                student_id = 0
        status = record.get("status")
        if not student_id or status not in {"Present", "Absent", "Late"}:
            return jsonify({"success": False, "message": "Resolve every included row to a student and valid attendance status."}), 400
        try:
            confidence = float(record.get("confidence", 100))
        except (TypeError, ValueError):
            return jsonify({"success": False, "message": "Attendance confidence must be a number from 0 to 100."}), 400
        if not 0 <= confidence <= 100:
            return jsonify({"success": False, "message": "Attendance confidence must be a number from 0 to 100."}), 400
        if confidence < 85 and record.get("verified") is not True:
            return jsonify({"success": False, "message": "Rows below 85% confidence must be manually verified before saving."}), 400
        student_ids.append(student_id)
    if len(student_ids) != len(set(student_ids)):
        return jsonify({"success": False, "message": "Duplicate student rows are not allowed."}), 400

    connection = get_db_connection()
    subject = connection.execute("SELECT id, subject_name FROM subjects WHERE id = ? AND class_id = ? AND active = 1", (subject_id, class_id)).fetchone()
    if subject is None:
        connection.close()
        return jsonify({"success": False, "message": "Subject is not in the selected class."}), 404
    roster = connection.execute(
        "SELECT id FROM students WHERE class_id = ? AND active = 1 AND id IN ({})".format(",".join("?" for _ in student_ids)),
        [class_id, *student_ids],
    ).fetchall()
    if len(roster) != len(student_ids):
        connection.close()
        return jsonify({"success": False, "message": "One or more students are not in the selected active class roster."}), 400
    for record in records:
        student_id = int(record["student_id"])
        status = record.get("status")
        confidence = float(record.get("confidence", 100)) / 100
        connection.execute(
            "INSERT INTO attendance (student_id, subject_id, date, status, method, confidence, recognized_by) VALUES (?, ?, ?, ?, 'SHEET', ?, ?) ON CONFLICT(student_id, subject_id, date) DO UPDATE SET status = excluded.status, method = 'SHEET', confidence = excluded.confidence, recognized_by = excluded.recognized_by, updated_at = CURRENT_TIMESTAMP",
            (student_id, subject_id, attendance_date, status, confidence, session.get("name", "FACULTY")),
        )
        create_attendance_notification(connection, int(student_id), subject["subject_name"], attendance_date, status)
        create_low_attendance_notification(connection, int(student_id))
    connection.commit()
    connection.close()
    return jsonify({"success": True, "saved": len(records), "message": "Attendance sheet confirmed.", "attendance_date": attendance_date})


# ==================================================
# LOGIN / AUTH
# ==================================================

@app.route("/")
def home():
    return render_template("login.html")


@app.route("/login", methods=["POST"])
def login():
    identifier = request.form.get("identifier", "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "").lower()

    if not identifier or not password or not role:
        return render_template("login.html", error="Please enter a valid login ID, password, and role.")

    connection = get_db_connection()

    if role == "student":
        student = connection.execute(
            """
            SELECT s.*, u.email, u.password_hash
            FROM students s
            LEFT JOIN users u ON u.email = s.email
            WHERE (s.student_uid = ? OR s.email = ?) AND s.active = 1
            LIMIT 1
            """,
            (identifier, identifier),
        ).fetchone()
        if student and student["password_hash"] and check_password_hash(student["password_hash"], password):
            session["user_id"] = student["id"]
            session["role"] = "student"
            session["email"] = student["email"]
            session["name"] = student["name"]
            session["student_id"] = student["id"]
            connection.close()
            return redirect(url_for("student_dashboard"))
        connection.close()
        return render_template("login.html", error="Student ID/email or password is incorrect.")

    user = connection.execute(
        "SELECT * FROM users WHERE email = ? AND role = ? LIMIT 1",
        (identifier, role),
    ).fetchone()
    if user and user["password_hash"] and check_password_hash(user["password_hash"], password):
        set_user_session(user, role)
        connection.close()
        if role in {"admin", "faculty"}:
            return redirect(url_for("class_selection"))
        return redirect(url_for("home"))

    connection.close()
    return render_template("login.html", error="Invalid email or password.")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


# ==================================================
# CLASS SELECTION
# ==================================================

@app.route("/class-selection", methods=["GET", "POST"])
def class_selection():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))

    connection = get_db_connection()
    academic_years = connection.execute("SELECT * FROM academic_years ORDER BY academic_year DESC").fetchall()
    if request.method == "POST":
        academic_year_id = request.form.get("academic_year_id")
        year = request.form.get("year")
        department = request.form.get("department")
        division = request.form.get("division")
        if not all([academic_year_id, year, department, division]):
            return render_template("class_selection.html", academic_years=academic_years, error="Please complete all class details.")

        class_row = connection.execute(
            "SELECT id FROM classes WHERE academic_year_id = ? AND year = ? AND department = ? AND division = ? LIMIT 1",
            (academic_year_id, year, department, division),
        ).fetchone()
        if class_row is None:
            connection.execute(
                "INSERT INTO classes (academic_year_id, year, department, division) VALUES (?, ?, ?, ?)",
                (academic_year_id, year, department, division),
            )
            class_row = connection.execute(
                "SELECT id FROM classes WHERE academic_year_id = ? AND year = ? AND department = ? AND division = ? LIMIT 1",
                (academic_year_id, year, department, division),
            ).fetchone()
        session["selected_class_id"] = class_row["id"]
        connection.close()
        return redirect(url_for("admin_dashboard"))

    connection.close()
    return render_template("class_selection.html", academic_years=academic_years)


# ==================================================
# DASHBOARDS
# ==================================================

@app.route("/admin-dashboard")
def admin_dashboard():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))

    connection = get_db_connection()
    class_row = connection.execute(
        "SELECT c.id, c.year, c.department, c.division, a.academic_year FROM classes c JOIN academic_years a ON a.id = c.academic_year_id WHERE c.id = ?",
        (class_id,),
    ).fetchone()
    total_students = connection.execute("SELECT COUNT(*) AS total FROM students WHERE class_id = ? AND active = 1", (class_id,)).fetchone()["total"]
    total_subjects = connection.execute("SELECT COUNT(*) AS total FROM subjects WHERE class_id = ? AND active = 1", (class_id,)).fetchone()["total"]
    today = date.today().isoformat()
    present_today = connection.execute("""
        SELECT COUNT(*) AS total
        FROM attendance a
        JOIN subjects s ON s.id = a.subject_id
        WHERE s.class_id = ? AND a.date = ? AND a.status = 'Present'
    """, (class_id, today)).fetchone()["total"]
    absent_today = connection.execute("""
        SELECT COUNT(*) AS total
        FROM attendance a
        JOIN subjects s ON s.id = a.subject_id
        WHERE s.class_id = ? AND a.date = ? AND a.status = 'Absent'
    """, (class_id, today)).fetchone()["total"]

    students = connection.execute("SELECT id FROM students WHERE class_id = ? AND active = 1", (class_id,)).fetchall()
    low_attendance = 0
    for student in students:
        stats = student_summary(student["id"])
        if stats["total"] and stats["percentage"] < 75:
            low_attendance += 1

    connection.close()
    return render_template(
        "admin_dashboard.html",
        class_row=class_row,
        total_students=total_students,
        total_subjects=total_subjects,
        present_today=present_today,
        absent_today=absent_today,
        low_attendance=low_attendance,
    )


@app.route("/student-dashboard")
def student_dashboard():
    if session.get("role") != "student":
        return redirect(url_for("home"))

    connection = get_db_connection()
    student = connection.execute("SELECT * FROM students WHERE id = ?", (session.get("student_id"),)).fetchone()
    if student is None:
        connection.close()
        return redirect(url_for("home"))

    class_row = connection.execute(
        "SELECT c.id, c.year, c.department, c.division, a.academic_year FROM classes c JOIN academic_years a ON a.id = c.academic_year_id WHERE c.id = ?",
        (student["class_id"],),
    ).fetchone()
    summary = student_summary(student["id"])
    subject_rows = connection.execute(
        """
        SELECT s.id, s.subject_code, s.subject_name, f.faculty_code
        FROM subjects s
        LEFT JOIN faculty f ON f.id = s.faculty_id
        WHERE s.class_id = ? AND s.active = 1
        ORDER BY s.subject_code
        """,
        (student["class_id"],),
    ).fetchall()
    subject_cards = []
    for row in subject_rows:
        counts = connection.execute(
            """
            SELECT COUNT(*) AS total,
                   SUM(CASE WHEN status = 'Present' THEN 1 ELSE 0 END) AS present,
                   SUM(CASE WHEN status = 'Absent' THEN 1 ELSE 0 END) AS absent,
                   SUM(CASE WHEN status = 'Late' THEN 1 ELSE 0 END) AS late
            FROM attendance WHERE student_id = ? AND subject_id = ?
            """,
            (student["id"], row["id"]),
        ).fetchone()
        total = counts["total"] or 0
        present = counts["present"] or 0
        absent = counts["absent"] or 0
        late = counts["late"] or 0
        percentage = round((present / total) * 100, 2) if total else 0
        subject_cards.append({
            "subject_name": row["subject_name"],
            "subject_code": row["subject_code"],
            "present": present,
            "absent": absent,
            "late": late,
            "total": total,
            "percentage": percentage,
            "faculty_code": row["faculty_code"],
        })

    notifications = connection.execute(
        "SELECT * FROM notifications WHERE student_id = ? ORDER BY created_at DESC LIMIT 5",
        (student["id"],),
    ).fetchall()
    unread_count = connection.execute("SELECT COUNT(*) AS total FROM notifications WHERE student_id = ? AND is_read = 0", (student["id"],)).fetchone()["total"]
    connection.close()

    return render_template(
        "student_dashboard.html",
        student=student,
        class_row=class_row,
        summary=summary,
        subject_cards=subject_cards,
        notifications=notifications,
        unread_count=unread_count,
    )


# ==================================================
# CRUD: Students and Subjects
# ==================================================

@app.route("/students")
def students():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))
    connection = get_db_connection()
    rows = connection.execute("SELECT * FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
    connection.close()
    return render_template("students.html", students=rows)


@app.route("/add-student", methods=["POST"])
def add_student():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))

    name = request.form.get("name", "").strip()
    roll_no = request.form.get("roll_no", "").strip()
    email = request.form.get("email", "").strip()
    phone = request.form.get("phone", "").strip()
    if not all([name, roll_no, email]):
        return redirect(url_for("students"))

    student_uid = generate_student_uid(class_id)
    connection = get_db_connection()
    try:
        connection.execute(
            "INSERT INTO students (student_uid, roll_no, name, email, phone, class_id, password_hash, active) VALUES (?, ?, ?, ?, ?, ?, ?, 1)",
            (student_uid, roll_no, name, email, phone, class_id, generate_password_hash("student123")),
        )
        connection.execute(
            "INSERT INTO users (name, email, password_hash, role) VALUES (?, ?, ?, 'student')",
            (name, email, generate_password_hash("student123")),
        )
        connection.commit()
    except Exception:
        connection.rollback()
    finally:
        connection.close()
    return redirect(url_for("students"))


@app.route("/edit-student/<int:student_id>")
def edit_student(student_id):
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    connection = get_db_connection()
    student = connection.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    connection.close()
    if not student:
        return redirect(url_for("students"))
    return render_template("edit_student.html", student=student)


@app.route("/update-student/<int:student_id>", methods=["POST"])
def update_student(student_id):
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    name = request.form.get("name", "").strip()
    roll_no = request.form.get("roll_no", "").strip()
    email = request.form.get("email", "").strip()
    phone = request.form.get("phone", "").strip()
    connection = get_db_connection()
    connection.execute(
        "UPDATE students SET name = ?, roll_no = ?, email = ?, phone = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (name, roll_no, email, phone, student_id),
    )
    connection.execute(
        "UPDATE users SET name = ?, email = ?, updated_at = CURRENT_TIMESTAMP WHERE email = (SELECT email FROM students WHERE id = ?) AND role = 'student'",
        (name, email, student_id),
    )
    connection.commit()
    connection.close()
    return redirect(url_for("students"))


@app.route("/delete-student/<int:student_id>", methods=["POST"])
def delete_student(student_id):
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    connection = get_db_connection()
    connection.execute("UPDATE students SET active = 0 WHERE id = ?", (student_id,))
    connection.commit()
    connection.close()
    return redirect(url_for("students"))


@app.route("/subjects")
def subjects():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))
    connection = get_db_connection()
    rows = connection.execute(
        "SELECT s.*, f.faculty_code FROM subjects s LEFT JOIN faculty f ON f.id = s.faculty_id WHERE s.class_id = ? AND s.active = 1 ORDER BY s.subject_code",
        (class_id,),
    ).fetchall()
    connection.close()
    return render_template("subjects.html", subjects=rows)


@app.route("/add-subject", methods=["POST"])
def add_subject():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))
    subject_name = request.form.get("subject_name", "").strip()
    subject_code = request.form.get("subject_code", "").strip()
    faculty_code = request.form.get("faculty_code", "").strip()
    connection = get_db_connection()
    faculty = connection.execute("SELECT id FROM faculty WHERE faculty_code = ? LIMIT 1", (faculty_code,)).fetchone() if faculty_code else None
    connection.execute(
        "INSERT INTO subjects (subject_code, subject_name, faculty_id, class_id, lecture_type, total_hours, active) VALUES (?, ?, ?, ?, 'THEORY', 30, 1)",
        (subject_code, subject_name, faculty["id"] if faculty else None, class_id),
    )
    connection.commit(); connection.close()
    return redirect(url_for("subjects"))


@app.route("/edit-subject/<int:subject_id>")
def edit_subject(subject_id):
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    connection = get_db_connection()
    subject = connection.execute("SELECT * FROM subjects WHERE id = ?", (subject_id,)).fetchone()
    connection.close()
    if not subject:
        return redirect(url_for("subjects"))
    return render_template("edit_subject.html", subject=subject)


@app.route("/update-subject/<int:subject_id>", methods=["POST"])
def update_subject(subject_id):
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    subject_name = request.form.get("subject_name", "").strip()
    subject_code = request.form.get("subject_code", "").strip()
    faculty_code = request.form.get("faculty_code", "").strip()
    connection = get_db_connection()
    faculty = connection.execute("SELECT id FROM faculty WHERE faculty_code = ? LIMIT 1", (faculty_code,)).fetchone() if faculty_code else None
    connection.execute(
        "UPDATE subjects SET subject_name = ?, subject_code = ?, faculty_id = ? WHERE id = ?",
        (subject_name, subject_code, faculty["id"] if faculty else None, subject_id),
    )
    connection.commit(); connection.close()
    return redirect(url_for("subjects"))


@app.route("/delete-subject/<int:subject_id>", methods=["POST"])
def delete_subject(subject_id):
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    connection = get_db_connection()
    connection.execute("UPDATE subjects SET active = 0 WHERE id = ?", (subject_id,))
    connection.commit(); connection.close()
    return redirect(url_for("subjects"))


# ==================================================
# ATTENDANCE & SHEETS
# ==================================================

@app.route("/manual-attendance", methods=["GET", "POST"])
def manual_attendance():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))

    connection = get_db_connection()
    if request.method == "POST":
        subject_id = request.form.get("subject_id")
        attendance_date = request.form.get("attendance_date")
        if not subject_id or not attendance_date:
            return redirect(url_for("manual_attendance"))
        students = connection.execute("SELECT * FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
        subject = connection.execute("SELECT * FROM subjects WHERE id = ?", (subject_id,)).fetchone()
        for student in students:
            status = request.form.get(f"status_{student['id']}")
            if not status:
                continue
            connection.execute(
                """
                INSERT INTO attendance (student_id, subject_id, date, status, method, confidence, recognized_by)
                VALUES (?, ?, ?, ?, 'MANUAL', 1.0, 'FACULTY')
                ON CONFLICT(student_id, subject_id, date)
                DO UPDATE SET status = excluded.status, method = 'MANUAL', updated_at = CURRENT_TIMESTAMP
                """,
                (student["id"], subject_id, attendance_date, status),
            )
            create_attendance_notification(connection, student["id"], subject["subject_name"], attendance_date, status)
            create_low_attendance_notification(connection, student["id"])
        connection.commit(); connection.close()
        return redirect(url_for("manual_attendance"))

    subjects = connection.execute("SELECT * FROM subjects WHERE class_id = ? AND active = 1 ORDER BY subject_code", (class_id,)).fetchall()
    students = connection.execute("SELECT * FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
    connection.close()
    return render_template("manual_attendance.html", subjects=subjects, students=students, today=date.today().isoformat())


@app.route("/generate-attendance-sheet", methods=["GET", "POST"])
def generate_attendance_sheet():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))
    connection = get_db_connection()
    subjects = connection.execute("SELECT * FROM subjects WHERE class_id = ? AND active = 1 ORDER BY subject_code", (class_id,)).fetchall()
    if request.method == "POST":
        subject_id = request.form.get("subject_id")
        attendance_date = request.form.get("attendance_date")
        subject = connection.execute("SELECT * FROM subjects WHERE id = ?", (subject_id,)).fetchone()
        students = connection.execute("SELECT * FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (class_id,)).fetchall()
        connection.close()
        return render_template("print_attendance_sheet.html", subject=subject, students=students, attendance_date=attendance_date)
    connection.close()
    return render_template("generate_attendance_sheet.html", subjects=subjects, today=date.today().isoformat())


@app.route("/upload-attendance-sheet", methods=["GET", "POST"])
def upload_attendance_sheet():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    class_id = session.get("selected_class_id")
    if not class_id:
        return redirect(url_for("class_selection"))
    connection = get_db_connection()
    subjects = connection.execute("SELECT * FROM subjects WHERE class_id = ? AND active = 1 ORDER BY subject_code", (class_id,)).fetchall()
    connection.close()
    if request.method == "POST":
        subject_id = request.form.get("subject_id")
        attendance_date = request.form.get("attendance_date")
        if not subject_id or not attendance_date or "attendance_file" not in request.files:
            return render_template("upload_attendance_sheet.html", subjects=subjects, error="Please select a subject, date, and file.")
        file = request.files["attendance_file"]
        if file.filename == "":
            return render_template("upload_attendance_sheet.html", subjects=subjects, error="Please choose a file.")
        if not allowed_file(file.filename):
            return render_template("upload_attendance_sheet.html", subjects=subjects, error="Only PNG, JPG, or JPEG images are allowed.")
        filename = secure_filename(file.filename)
        unique_name = f"attendance_{uuid.uuid4()}_{filename}"
        file_path = os.path.join(app.config["UPLOAD_FOLDER"], unique_name)
        file.save(file_path)
        return render_template("upload_attendance_sheet.html", subjects=subjects, uploaded_file=unique_name, selected_subject_id=subject_id, selected_date=attendance_date, message="Attendance sheet uploaded successfully.")
    return render_template("upload_attendance_sheet.html", subjects=subjects)


@app.route("/uploads/<filename>")
def uploaded_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


@app.route("/process-attendance-sheet", methods=["POST"])
def process_attendance_sheet():
    uploaded_file_name = request.form.get("uploaded_file")
    subject_id = request.form.get("subject_id")
    attendance_date = request.form.get("attendance_date")
    if not uploaded_file_name:
        return "Uploaded file is missing."
    file_path = os.path.join(app.config["UPLOAD_FOLDER"], os.path.basename(uploaded_file_name))
    if not os.path.exists(file_path):
        return "Uploaded file not found."
    image = cv2.imread(file_path)
    if image is None:
        return "Could not read uploaded image."

    connection = get_db_connection()
    subject = connection.execute("SELECT * FROM subjects WHERE id = ?", (subject_id,)).fetchone()
    students = connection.execute("SELECT * FROM students WHERE class_id = ? AND active = 1 ORDER BY roll_no", (session.get("selected_class_id"),)).fetchall()
    connection.close()

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    _, binary = cv2.threshold(gray, 180, 255, cv2.THRESH_BINARY_INV)
    contours, _ = cv2.findContours(binary, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    checkbox_centers = []
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        if 15 <= w <= 100 and 15 <= h <= 100 and 0.6 <= w / h <= 1.4:
            checkbox_centers.append((x + w // 2, y + h // 2, w, h))

    results = []
    if students:
        image_height = image.shape[0]
        table_top = int(image_height * 0.30)
        table_bottom = int(image_height * 0.75)
        row_height = (table_bottom - table_top) / len(students)
        image_width = image.shape[1]
        present_x = int(image_width * 0.78)
        absent_x = int(image_width * 0.92)
        for index, student in enumerate(students):
            row_center = int(table_top + (row_height * (index + 0.5)))
            present_found = False
            absent_found = False
            for center_x, center_y, _, _ in checkbox_centers:
                vertical_distance = abs(center_y - row_center)
                if vertical_distance > row_height * 0.35:
                    continue
                if abs(center_x - present_x) < image_width * 0.08:
                    present_found = True
                if abs(center_x - absent_x) < image_width * 0.08:
                    absent_found = True
            if present_found and not absent_found:
                status = "Present"
                confidence = "High"
            elif absent_found and not present_found:
                status = "Absent"
                confidence = "High"
            else:
                status = "Review"
                confidence = "Low"
            results.append({"student_id": student["id"], "roll_no": student["roll_no"], "name": student["name"], "status": status, "confidence": confidence})

    return render_template("attendance_sheet_result.html", subject=subject, attendance_date=attendance_date, uploaded_file=uploaded_file_name, results=results, has_review=any(r["status"] == "Review" for r in results))


@app.route("/confirm-sheet-attendance", methods=["POST"])
def confirm_sheet_attendance():
    subject_id = request.form.get("subject_id")
    attendance_date = request.form.get("attendance_date")
    student_ids = request.form.getlist("student_id")
    connection = get_db_connection()
    subject = connection.execute("SELECT subject_name FROM subjects WHERE id = ?", (subject_id,)).fetchone()
    for student_id in student_ids:
        status = request.form.get(f"status_{student_id}")
        if status not in {"Present", "Absent"}:
            continue
        connection.execute(
            """
            INSERT INTO attendance (student_id, subject_id, date, status, method, confidence)
            VALUES (?, ?, ?, ?, 'SHEET', 1.0)
            ON CONFLICT(student_id, subject_id, date)
            DO UPDATE SET status = excluded.status, method = 'SHEET', updated_at = CURRENT_TIMESTAMP
            """,
            (student_id, subject_id, attendance_date, status),
        )
        create_attendance_notification(connection, int(student_id), subject["subject_name"], attendance_date, status)
        create_low_attendance_notification(connection, int(student_id))
    connection.commit(); connection.close()
    return render_template("attendance_saved.html", attendance_date=attendance_date)


# ==================================================
# NOTIFICATIONS
# ==================================================

@app.route("/notifications")
def notifications():
    if session.get("role") != "student":
        return redirect(url_for("home"))
    student = get_db_connection().execute("SELECT * FROM students WHERE id = ?", (session.get("student_id"),)).fetchone()
    rows = get_db_connection().execute("SELECT * FROM notifications WHERE student_id = ? ORDER BY created_at DESC", (student["id"],)).fetchall()
    get_db_connection().execute("UPDATE notifications SET is_read = 1 WHERE student_id = ?", (student["id"],))
    return render_template("notifications.html", student=student, notifications=rows)


@app.route("/admin-notifications", methods=["GET", "POST"])
def admin_notifications():
    if session.get("role") not in {"admin", "faculty"}:
        return redirect(url_for("home"))
    connection = get_db_connection()
    if request.method == "POST":
        title = request.form.get("title", "").strip()
        message = request.form.get("message", "").strip()
        if title and message:
            students = connection.execute("SELECT id FROM students WHERE class_id = ? AND active = 1", (session.get("selected_class_id"),)).fetchall()
            for student in students:
                create_notification(student["id"], session.get("selected_class_id"), None, title, message, "ANNOUNCEMENT", "HIGH")
        connection.commit(); connection.close();
        return render_template("admin_notifications.html", notifications=connection.execute("SELECT * FROM notifications WHERE type = 'ANNOUNCEMENT' ORDER BY created_at DESC LIMIT 20").fetchall(), message="Announcement sent successfully.")
    rows = connection.execute("SELECT * FROM notifications WHERE type = 'ANNOUNCEMENT' ORDER BY created_at DESC LIMIT 20").fetchall()
    connection.close();
    return render_template("admin_notifications.html", notifications=rows)


# ==================================================
# RUN APP
# ==================================================

if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1")